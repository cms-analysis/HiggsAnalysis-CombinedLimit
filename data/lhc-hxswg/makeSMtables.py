import logging
import os, sys

from openpyxl import load_workbook


stdHeading = ("mH_GeV", "XS_pb", "Sca_Hi", "Sca_Lo", "Pdf_alpha_s", "Pdf", "alpha_s")
xsecGroups = {
    "ggH": {"col": "A", "heading": stdHeading},
    "VBF": {"col": "I", "heading": stdHeading},
    "WH": {"col": "Q", "heading": stdHeading + ("XS_WplusH_pb", "XS_WminusH_pb")},
    "ZH": {"col": "AA", "heading": stdHeading + ("XS_ggZH_pb",)},
    "ttH": {"col": "AJ", "heading": stdHeading},
    "bbH": {
        "col": "AR",
        "heading": ("mH_GeV", "XS_pb", "Sca_Pdf_mb_mub_Hi", "Sca_Pdf_mb_mub_Lo"),
    },
    "tH_tchan": {"col": "AZ", "heading": stdHeading + ("XS_tH_pb", "XS_tbarH_pb")},
    "tH_schan": {"col": "BJ", "heading": stdHeading + ("XS_tH_pb", "XS_tbarH_pb")},
    # 'total':  {'col':'BT', 'heading':('XS_pb',)},
    "WminusH_lv": {"col": "BX", "heading": stdHeading + ("XS_gamma_pb",)},
    "WplusH_lv": {"col": "CG", "heading": stdHeading + ("XS_gamma_pb",)},
    "ZH_ll": {"col": "CP", "heading": stdHeading + ("XS_ggZH_pb", "XS_gamma_pb")},
    "ZH_vv": {"col": "CZ", "heading": stdHeading + ("XS_ggZH_pb", "XS_gamma_pb")},
    "VBF_qqH_schan": {"col": "DJ", "heading": ("mH_GeV", "XS_pb")},
}

reducedHeading = ("mH_GeV", "XS_pb", "Sca_Hi", "Sca_Lo", "Pdf_alpha_s")
xsecGroupsBSM = {
    "ggH": {"col": "A", "heading": stdHeading + ("1_plus_dEW",)},
    "VBF": {"col": "J", "heading": stdHeading},
    "WH": {"col": "S", "heading": reducedHeading},
    "ZH": {"col": "AD", "heading": reducedHeading},
    "bbH": {
        "col": "AW",
        "heading": ("mH_GeV", "XS_pb", "Sca_Pdf_mb_mub_Hi", "Sca_Pdf_mb_mub_Lo"),
    },
    "WminusH": {"col": "CF", "heading": reducedHeading},
    "WplusH": {"col": "CO", "heading": reducedHeading},
}


#mH_GeV  XS_pb       Scale_pos	Scale_neg Gauss PDF_plus_alpha_s  PDF alpha_S
stdHeadingYR5 = ("mH_GeV", "XS_pb", "Sca_Hi", "Sca_Lo", "Pdf_alpha_S", "Total_Hi", "Total_Lo")
xsecGroupsYR5 = {
    "ggH": { "col": "A", "heading": ("mH_GeV","XS_pb", "Sca_Hi", "Sca_Lo", "Pdf_TH", "Gauss", "Pdf_alpha_s", "Total_Hi","Total_Lo","Total_Gauss")},
    "VBF": {"col": "L", "heading": stdHeadingYR5},
    "WH" : {"col": "T", "heading": stdHeadingYR5 + ("XS_WplusH_pb", "XS_WminusH_pb")},
    #"WplusH": {"col": "AA", "heading": ("XS_pb",)},
    #"WminusH": {"col": "AB", "heading": ("XS_pb",)},
    "ZH" : {"col": "AD", "heading": stdHeadingYR5 + ("XS_ggZH_pb",)}, 
    #"ggZH" : {"col": "AK", "heading": ("XS_pb")},
    "ttH" : {"col": "AM", "heading": stdHeadingYR5},
    "bbH" : {"col": "AU", "heading": stdHeadingYR5},
    "tH_tchan" : {"col": "BM", "heading": stdHeadingYR5 + ("XS_tH_pb", "XS_tbarH_pb")},
    "tH_schan" : {"col": "BC", "heading": stdHeadingYR5 + ("XS_tH_pb", "XS_tbarH_pb")},
    "tH_Wassoc" : {"col": "BW", "heading": stdHeadingYR5 },
}

specsYR5 = { key: {"rows": (6, 19), "groups": xsecGroupsYR5} for key in ["7 TeV", "8 TeV", "13 TeV", "13.6 TeV","14 TeV"] }    


specs = {

    "YR4 SM 7TeV": {
        "rows": (6, 43),
        "groups": xsecGroups,
    },
    "YR4 SM 8TeV": {
        "rows": (6, 43),
        "groups": xsecGroups,
    },
    "YR4 SM 13TeV": {
        "rows": (6, 43),
        "groups": xsecGroups,
    },
    "YR4 SM 14TeV": {
        "rows": (6, 43),
        "groups": xsecGroups,
    },
    "YR4 BSM 7TeV": {
        "rows": (6, 119),
        "groups": xsecGroupsBSM,
    },
    "YR4 BSM 8TeV": {
        "rows": (6, 119),
        "groups": xsecGroupsBSM,
    },
    "YR4 BSM 13TeV": {
        "rows": (6, 119),
        "groups": xsecGroupsBSM,
    },
    "YR4 BSM 14TeV": {
        "rows": (6, 119),
        "groups": xsecGroupsBSM,
    },
    "YR4 SM BR": {
        "rows": (7, 44),
        "groups": {
            "BR1": {
                "col": "A",
                "heading": (
                    "mH_GeV",
                    "H_bb",
                    "THU_Hi",
                    "THU_Lo",
                    "PU_mq_Hi",
                    "PU_mq_Lo",
                    "PU_as_Hi",
                    "PU_as_Lo",
                    "H_tautau",
                    "THU_Hi",
                    "THU_Lo",
                    "PU_mq_Hi",
                    "PU_mq_Lo",
                    "PU_as_Hi",
                    "PU_as_Lo",
                    "H_mumu",
                    "THU_Hi",
                    "THU_Lo",
                    "PU_mq_Hi",
                    "PU_mq_Lo",
                    "PU_as_Hi",
                    "PU_as_Lo",
                    "H_ccbar",
                    "THU_Hi",
                    "THU_Lo",
                    "PU_mq_Hi",
                    "PU_mq_Lo",
                    "PU_as_Hi",
                    "PU_as_Lo",
                ),
            },
            "BR": {
                "col": "AS",
                "heading": (
                    "mH_GeV",
                    "H_gg",
                    "THU_Hi",
                    "THU_Lo",
                    "PU_mq_Hi",
                    "PU_mq_Lo",
                    "PU_as_Hi",
                    "PU_as_Lo",
                    "H_gamgam",
                    "THU_Hi",
                    "THU_Lo",
                    "PU_mq_Hi",
                    "PU_mq_Lo",
                    "PU_as_Hi",
                    "PU_as_Lo",
                    "H_Zgam",
                    "THU_Hi",
                    "THU_Lo",
                    "PU_mq_Hi",
                    "PU_mq_Lo",
                    "PU_as_Hi",
                    "PU_as_Lo",
                    "H_WW",
                    "THU_Hi",
                    "THU_Lo",
                    "PU_mq_Hi",
                    "PU_mq_Lo",
                    "PU_as_Hi",
                    "PU_as_Lo",
                    "H_ZZ",
                    "THU_Hi",
                    "THU_Lo",
                    "PU_mq_Hi",
                    "PU_mq_Lo",
                    "PU_as_Hi",
                    "PU_as_Lo",
                    "Total_Width_GeV",
                    "THU_Hi",
                    "THU_Lo",
                    "PU_mq_Hi",
                    "PU_mq_Lo",
                    "PU_as_Hi",
                    "PU_as_Lo",
                ),
            },
            "BR2": {
                "col": "CK",
                "heading": (
                    "mH_GeV",
                    "H_llll_emt",
                    "H_llll_em",
                    "H_eeee",
                    "H_eemm",
                    "H_llvv_emt",
                    "H_evev",
                    "H_llqq_emt",
                    "H_llqq_em",
                    "H_lvqq_em",
                    "H_vvqq",
                    "H_qqqq",
                    "H_ffff",
                    "DBR",
                ),
            },
        },
    },
}

morespecs = {
    #'sm/xs/7TeV/7TeV-ggH.txt'
}

## Add YR5 as well
specs = {**specs, **specsYR5}

def find_starting_points(o):
    import string

    def find_col_by_substring(vals, cols, offset, needle):
        idx = next((i for i, val in enumerate(vals) if needle in str(val)), None)
        if idx is None:
            return None
        return cols[idx + offset]

    cols = [ x+c for x in ["","A","B","C"] for c in string.ascii_uppercase ] 
    d = {
        "What": {"col": "A", "heading": cols }
    }
    specs = { key: {"rows": (2,2), "groups": d} for key in ["7 TeV", "8 TeV", "13 TeV", "13.6 TeV","14 TeV"] }    
    f = open_workbook(o.input)
    for s in f.sheets():
        spec = specs.get(s.name)
        if spec is None:
            logging.info("Skipping sheet [%s]: I do not have parsing rules for it.", s.name)
            continue
        logging.info("Processing sheet [" + s.name + "]")
        for group, props in spec["groups"].items():
            logging.info("Processing [" + group + "] in [" + s.name + "]")
            # open output
            # dump heading
            heading = props["heading"]
            startRow, endRow = spec["rows"]
            startings={}
            for r in range(startRow - 1, endRow):
                offset = col2num(props["col"]) - 1
                vals = s.row_values(r)[offset : offset + len(heading)]
                print(vals)
                for key in ['ggF', 'VBF', 'WH', 'ZH', 'ttH', 'bbH', 'tH (s-ch', 'tH (t-ch', 'tH (W-ass']:
                    col = find_col_by_substring(vals, cols, offset, key)
                    if col is not None:
                        startings[key] = col
        print('------------------------------------------------------------')
        print(s.name, startings)
        print('------------------------------------------------------------')
    print("==============================================================")


# import prettytable
def print_table(table):
    col_width = [max(len(x) for x in col) for col in zip(*table)]
    for line in table:
        print("  ".join("{:{}}".format(x, col_width[i]) for i, x in enumerate(line)))


# Based on http://stackoverflow.com/a/12640614/665025
def col2num(col_str):
    """Convert base26 column string to number."""
    expn = 0
    col_num = 0
    for char in reversed(col_str):
        col_num += (ord(char) - ord("A") + 1) * (26**expn)
        expn += 1
    return col_num


# import urllib2
# response = urllib2.urlopen('https://twiki.cern.ch/twiki/pub/LHCPhysics/CERNYellowReportPageAt8TeV/Higgs_XSBR_YR3.xlsx')
# f = xlrd.open_workbook(file_contents=response.read())


class OpenPyXLSheet(object):
    def __init__(self, ws):
        self.name = ws.title
        # values_only=True returns plain Python values; normalize None to ''
        self._rows = [list(row) for row in ws.iter_rows(values_only=True)]

    def row_values(self, idx):
        row = self._rows[idx]
        return ["" if v is None else v for v in row]


class OpenPyXLBook(object):
    def __init__(self, wb):
        self._sheets = [OpenPyXLSheet(ws) for ws in wb.worksheets]

    def sheets(self):
        return self._sheets


def open_workbook(path):
    if path.lower().endswith(".xlsx"):
        wb = load_workbook(filename=path, data_only=True, read_only=True)
        return OpenPyXLBook(wb)

    if xlrd is None:
        raise RuntimeError("Legacy .xls input requires xlrd, but xlrd is not installed.")

    return xlrd.open_workbook(path)


def formatval(v):
    try:
        # This is to remove pesky unicode symbols like \pm
        v = float(v.encode("ascii", "ignore"))
    except AttributeError:
        pass
    return "{:+}".format(v)


def main(o):
    f = open_workbook(o.input)
    for s in f.sheets():
        try:
            spec = specs[s.name]
        except KeyError:
            logging.info("Skipping sheet [%s]: I do not have parsing rules for it.", s.name)
            continue
        logging.info("Processing sheet [" + s.name + "]")
        for group, props in spec["groups"].items():
            table = []
            logging.info("Processing [" + group + "] in [" + s.name + "]")
            # open output
            # dump heading
            heading = props["heading"]
            table.append(heading)
            startRow, endRow = spec["rows"]
            for r in range(startRow - 1, endRow):
                offset = col2num(props["col"]) - 1
                vals = s.row_values(r)[offset : offset + len(heading)]
                if set(vals[1:]) == set(("",)):
                    continue
                try:
                    table.append(list(map(formatval, vals)))
                except ValueError:
                    print("Could not parse the followig tuple: ")
                    print(vals)
                    raise

            print_table(table)

            if o.write:
                # sm/xs/13TeV/13TeV-tHW.txt
                sqrtS = s.name.replace(" ", "").replace(".", "p")
                #13TeV-WH.txt  13TeV-ZH.txt  13TeV-bbH.txt  13TeV-ggH-NNLO-NLL.txt  13TeV-ggH.txt  13TeV-ggZH.txt  13TeV-tHW.txt  13TeV-tHq.txt  13TeV-ttH.txt  13TeV-vbfH.txt 
                procname = group[:]
                if group == 'VBF': procname = 'vbfH'
                if group == 'ZH': procname = 'ZH'
                if group == 'tH_Wassoc': procname = 'tHW'

                if True: ## thisi is the default behaviour
                    os.makedirs(f"sm/xs/{sqrtS}", exist_ok=True)
                    outname = f"sm/xs/{sqrtS}/{sqrtS}-{procname}.txt"
                    with open(outname, "w") as out:
                        for line in table:
                            out.write("  ".join(line) + "\n")
                        if group == 'tH_Wassoc' and sqrtS in ["7TeV", "8TeV"]:
                            ## add 120 and 130 as duplicate of 125
                            for mh in [120, 130]:
                                line2 = list(table[1][:])
                                line2[0] = str(mh)
                                line2[1] = table[1][1] ## tHW XS_pb
                                out.write("  ".join(line2) + "\n")

                if group == 'ZH':
                    outname = f"sm/xs/{sqrtS}/{sqrtS}-ggZH.txt"
                    with open(outname, "w") as out:
                        heading = ["mH_GeV","XS_pb" , "Sca_Hi"  , "Sca_Lo"  , "Pdf_alpha_S" ]
                        for idx,line in enumerate(table):
                            if idx == 0:
                                out.write("  ".join(heading) + "\n")
                            else:
                                line2 = list(line[:5])
                                line2[1] = line[7] ## ggZH XS_pb
                                out.write("  ".join(line2) + "\n")
                    #outname = f"sm/xs/{sqrtS}/{sqrtS}-qqZH.txt"
                    #with open(outname, "w") as out:
                    #    heading = ["mH_GeV","XS_pb" , "Sca_Hi"  , "Sca_Lo"  , "Pdf_alpha_S" ]
                    #    for idx,line in enumerate(table):
                    #        if idx == 0:
                    #            out.write("  ".join(heading) + "\n")
                    #        else:
                    #            line2 = list(line[:5])
                    #            line2[1] = formatval( float(line[1])-float(line[7]) ) ## ZH - ggZH XS_pb
                    #            out.write("  ".join(line2) + "\n")

                #13TeV-WH.txt  13TeV-ZH.txt  13TeV-bbH.txt  13TeV-ggH-NNLO-NLL.txt  13TeV-ggH.txt  13TeV-ggZH.txt  13TeV-tHW.txt  13TeV-tHq.txt  13TeV-ttH.txt  13TeV-vbfH.txt  README.txt



if __name__ == "__main__":
    from optparse import OptionParser

    parser = OptionParser(usage="%prog -i FILE.xls[x]", version="%prog 3.141")

    parser.add_option("-i", "--input", type="string", help="HXSWG XSBR file", metavar="FILE")
    parser.add_option(
        "-l",
        "--log",
        default="INFO",
        metavar="LEVEL",
        help="Set the minimum logging level.",
    )

    parser.add_option("-x","--find-starting-points", action="store_true", help="Find the starting points for the tables, and exit.", default=False)
    parser.add_option("-w","--write", action="store_true", help="Write",default=False)

    o, args = parser.parse_args()

    if not o.input:
        parser.error("Please specify an input Excel file from the LHC HXSWG.")

    logging.basicConfig(level=getattr(logging, o.log.upper()))

    logging.debug("%s" % str(o))

    if o.find_starting_points:
        print("Finding starting points for the tables...")
        sys.exit(find_starting_points(o))

    sys.exit(main(o))
